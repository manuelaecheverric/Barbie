"""
Paso 1 de la carga: leer el archivo y extraer su texto.

Cada lector devuelve una lista de Document de LangChain (texto + metadatos).
Formatos: PDF, TXT, MD, DOCX, CSV.
"""

import csv
import io
from pathlib import Path

from langchain_core.documents import Document

FORMATOS_SOPORTADOS = {".pdf", ".txt", ".md", ".docx", ".csv"}


def _leer_texto_plano(ruta: Path) -> str:
    """Lee un archivo de texto. Prueba UTF-8 y, si falla, la codificación de Windows."""
    datos = ruta.read_bytes()
    for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return datos.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return datos.decode("utf-8", errors="replace")


def _leer_txt_o_md(ruta: Path) -> list[Document]:
    return [Document(page_content=_leer_texto_plano(ruta))]


def _leer_pdf(ruta: Path) -> list[Document]:
    from pypdf import PdfReader

    lector = PdfReader(str(ruta))
    paginas = []
    for numero, pagina in enumerate(lector.pages, start=1):
        texto = (pagina.extract_text() or "").strip()
        if texto:
            paginas.append(Document(page_content=texto, metadata={"pagina": numero}))
    if not paginas:
        raise ValueError(
            "El PDF no tiene texto que se pueda extraer (puede ser un PDF escaneado)."
        )
    return paginas


def _leer_docx(ruta: Path) -> list[Document]:
    """Lee un Word. Los títulos se convierten a Markdown (#, ##, ###)
    para que el fragmentador los use igual que en los archivos .md."""
    from docx import Document as WordDocument

    word = WordDocument(str(ruta))
    lineas = []
    for parrafo in word.paragraphs:
        texto = parrafo.text.strip()
        if not texto:
            continue
        estilo = (parrafo.style.name or "").lower()
        if estilo == "title" or estilo == "título":
            lineas.append(f"# {texto}")
        elif estilo.startswith(("heading", "título")):
            nivel = "".join(c for c in estilo if c.isdigit()) or "1"
            lineas.append(f"{'#' * min(int(nivel) + 1, 3)} {texto}")
        else:
            lineas.append(texto)
        lineas.append("")
    for tabla in word.tables:
        for fila in tabla.rows:
            celdas = [celda.text.strip() for celda in fila.cells]
            lineas.append(" | ".join(celdas))
        lineas.append("")
    return [Document(page_content="\n".join(lineas))]


def _leer_csv(ruta: Path) -> list[Document]:
    """Cada fila se convierte en una línea 'columna: valor | columna: valor'."""
    texto = _leer_texto_plano(ruta)
    try:
        dialecto = csv.Sniffer().sniff(texto[:2048], delimiters=",;\t")
    except csv.Error:
        dialecto = csv.excel
    filas = csv.DictReader(io.StringIO(texto), dialect=dialecto)
    lineas = []
    for fila in filas:
        partes = [f"{col}: {val}" for col, val in fila.items() if col and val]
        if partes:
            lineas.append(" | ".join(partes))
    return [Document(page_content="\n".join(lineas))]


_LECTORES = {
    ".txt": _leer_txt_o_md,
    ".md": _leer_txt_o_md,
    ".pdf": _leer_pdf,
    ".docx": _leer_docx,
    ".csv": _leer_csv,
}


def leer_archivo(file_path: str, extension: str) -> list[Document]:
    extension = extension.lower()
    if extension not in _LECTORES:
        raise ValueError(f"Formato no soportado: {extension}")
    return _LECTORES[extension](Path(file_path))
